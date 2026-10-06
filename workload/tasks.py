"""Deterministic, bounded work; document packages are imported only for that scenario."""

import argparse
import json
import os
import subprocess
import time
import urllib.request
from pathlib import Path

ROOT = Path("/workspace")
ASSETS = os.environ.get("BENCH_ASSETS", "http://bench-assets.ax-ram-bench.svc.cluster.local")
TESTS = [
    "tests/test_get_request_body.py",
    "tests/test_schema_extra_examples.py",
    "tests/test_tutorial/test_first_steps/test_tutorial001.py",
]


def command(argv, cwd=None):
    subprocess.run(argv, cwd=cwd, check=True, timeout=180)


def documents():
    import pandas as pd
    from docx import Document
    from openpyxl import load_workbook

    dst = ROOT / "documents"
    dst.mkdir(exist_ok=True)
    df = pd.DataFrame(
        {
            "id": range(10000),
            "group": [i % 20 for i in range(10000)],
            "amount": [(i * 17) % 1000 for i in range(10000)],
        }
    )
    df.to_csv(dst / "input.tmp", index=False)
    (dst / "input.tmp").replace(dst / "input.csv")
    loaded = pd.read_csv(dst / "input.csv")
    totals = loaded.groupby("group")["amount"].sum()
    with pd.ExcelWriter(dst / "sample.xlsx", engine="openpyxl") as writer:
        loaded.to_excel(writer, index=False, sheet_name="Data")
        totals.to_excel(writer, sheet_name="Summary")
    book = load_workbook(dst / "sample.xlsx", read_only=True)
    assert book["Data"].max_row == 10001 and book["Data"]["A2"].value == 0
    book.close()
    doc = Document()
    doc.add_heading("Quarterly document benchmark", 0)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Group"
    table.rows[0].cells[1].text = "Total"
    for k, v in totals.items():
        cells = table.add_row().cells
        cells[0].text = str(k)
        cells[1].text = str(v)
    doc.add_paragraph("Deterministic Python, filesystem, spreadsheet and document work.")
    doc.save(dst / "sample.docx")
    reopened = Document(dst / "sample.docx")
    assert len(reopened.tables[0].rows) == 21
    assert int(totals.sum()) == int(loaded.amount.sum())
    return {"rows": len(loaded), "xlsx_valid": True, "docx_valid": True}


def dev_prepare():
    repo = ROOT / "fastapi"
    command(["git", "clone", ASSETS + "/fastapi.git", str(repo)])
    command(["git", "checkout", "bb8c2a64981d3e806575a445115f29eddf014c77"], repo)
    wheels = ROOT / "wheels"
    wheels.mkdir(exist_ok=True)
    names = json.load(urllib.request.urlopen(ASSETS + "/wheels.json"))
    for name in names:
        urllib.request.urlretrieve(ASSETS + "/wheels/" + name, wheels / name)
    command(["python", "-m", "venv", str(ROOT / "venv")])
    py = str(ROOT / "venv/bin/python")
    command(
        [
            py,
            "-m",
            "pip",
            "install",
            "--no-index",
            "--find-links",
            str(wheels),
            "--require-hashes",
            "-r",
            "/opt/bench/dev-requirements.lock",
        ]
    )
    # Import directly from the real clone, and build its distributable each cycle.
    (ROOT / "dev-ready").write_text("bb8c2a64981d3e806575a445115f29eddf014c77")
    return developer()


def developer():
    py = str(ROOT / "venv/bin/python")
    command([py, "-m", "pytest", "-q", "-o", "addopts=", *TESTS], ROOT / "fastapi")
    command(
        [py, "-m", "build", "--wheel", "--no-isolation", "--outdir", str(ROOT / "dist")],
        ROOT / "fastapi",
    )
    assert list((ROOT / "dist").glob("fastapi-*.whl"))
    return {"tests_passed": True, "wheel_built": True}


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("scenario", choices=["documents", "dev-prepare", "developer"])
    p.add_argument("--seconds", type=int, default=0)
    args = p.parse_args()
    end = time.monotonic() + args.seconds
    cycle = 0
    while True:
        start = time.time()
        result = {"documents": documents, "dev-prepare": dev_prepare, "developer": developer}[
            args.scenario
        ]()
        cycle += 1
        print(
            json.dumps({"cycle": cycle, "start": start, "end": time.time(), "result": result}),
            flush=True,
        )
        if time.monotonic() >= end:
            break
        time.sleep(0.5)
